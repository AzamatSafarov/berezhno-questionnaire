import json,os,threading,tempfile
from pathlib import Path
from functools import partial
from http.server import ThreadingHTTPServer,SimpleHTTPRequestHandler
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
class Quiet(SimpleHTTPRequestHandler):
 def log_message(self,*args):pass
server=ThreadingHTTPServer(('127.0.0.1',0),partial(Quiet,directory=str(ROOT/'public')))
thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
url=f'http://127.0.0.1:{server.server_port}'
results=[]
try:
 with sync_playwright()as p,tempfile.TemporaryDirectory()as tmp:
  try:b=p.chromium.launch(channel='msedge',headless=True)
  except Exception:b=p.chromium.launch(headless=True)
  c=b.new_context(accept_downloads=True);q=c.new_page();errors=[];q.on('pageerror',lambda e:errors.append(str(e)))
  q.goto(url);q.wait_for_selector('#start');assert q.locator('.method-tile').count()==10;q.locator('#consent').check();q.locator('#start').click();q.locator('#next').click();assert q.locator('#validation').inner_text();assert q.locator('[data-skip]').count()==0
  pages=0
  while q.locator('#next').count():
   for card in q.locator('.question').all():
    radio=card.locator('input[type=radio]')
    if radio.count():radio.last.check()
    else:card.locator('textarea.open-answer').fill('Нет')
   q.locator('#next').click();pages+=1;assert pages<70
  assert '240 из 240' in q.locator('.summary-intro').inner_text()
  with q.expect_download()as d:q.locator('#download-json').click()
  path=Path(tmp)/'synthetic.json';d.value.save_as(path);data=json.loads(path.read_text(encoding='utf-8'));assert data['answered']==240 and data['unanswered']==0 and data['skipped']==0;assert list(data['scores'])==['MSPSS'];assert data['scores']['MSPSS']['total']==84
  q.locator('#finish-local').click();assert q.locator('#finish-local').is_disabled();assert 'Спасибо за участие' in q.locator('.summary h1').inner_text();assert 'на этом устройстве' in q.locator('.summary-intro').inner_text();assert q.evaluate('JSON.parse(localStorage.getItem(KEY)).state.localCompleted')
  results.append('full 240-item local flow, incomplete blocking, MSPSS=84, finish storage read-back')
  q.goto(url+'/admin.html');q.wait_for_function('()=>typeof buildResearchWorkbook==="function"');fixture=q.evaluate('()=>{const answers={};QUESTIONNAIRE_DATA.forEach(m=>m.questions.forEach(q=>answers[q.id]=q.open?"Нет":(q.options||m.options)[0].value));return {user_code:"SYNTHETIC",completed_at:"2026-10-12",is_test:false,answers,notes:{},answered_count:240,skipped_count:0}}')
  incomplete={**fixture,'answers':{**fixture['answers'],'MSPSS-1':'skip'}}
  with q.expect_download()as d:q.evaluate('(rows)=>XLSX.writeFile(buildResearchWorkbook(rows),"synthetic.xlsx")',[fixture,incomplete])
  path=Path(tmp)/'synthetic.xlsx';d.value.save_as(path);w=load_workbook(path);assert w['Участницы'].max_row==2;assert w['Ответы'].max_row==241;assert len(w.sheetnames)==5;results.append('five-sheet admin Excel excludes incomplete record')
  assert not errors,errors;c.close()
  for width in [320,360,390,430,768,1440]:
   c=b.new_context(viewport={'width':width,'height':844},is_mobile=width<768,has_touch=width<768);q=c.new_page();q.goto(url);q.wait_for_selector('#start');assert q.evaluate('document.documentElement.scrollWidth<=innerWidth');q.locator('#consent').check();q.locator('#start').click()
   for i in range(10):
    q.locator(f'[data-method="{i}"]').click();assert q.evaluate('document.documentElement.scrollWidth<=innerWidth'),(width,i)
    if width<600:assert q.locator('.option').first.bounding_box()['height']>=48
    radio=q.locator('input[type=radio]').first;radio.focus();q.keyboard.press('Space');assert radio.is_checked()
   q.evaluate('()=>{state.screen="summary";render()}');q.wait_for_selector('#finish-local');assert q.locator('#finish-local').is_disabled();assert q.evaluate('document.documentElement.scrollWidth<=innerWidth');q.goto(url+'/admin.html');q.wait_for_selector('#send-code');assert q.evaluate('document.documentElement.scrollWidth<=innerWidth');c.close()
  results.append('320/360/390/430/768/1440px: ten methods, keyboard radios, summary, admin login')
  b.close()
 print(json.dumps({'status':'passed','checks':results,'not_tested':['real email login','new Supabase migration replay','live invited server completion']},ensure_ascii=False,indent=2))
finally:server.shutdown();server.server_close()
